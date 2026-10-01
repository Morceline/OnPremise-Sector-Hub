"""
services/document.py
----------------------
Leitor seguro de arquivos locais (PDF, DOCX, TXT), respeitando a
allowlist/blocklist de caminhos definida pelo gestor do setor (RF01).

Mudanças:
- Adicionado limite de tamanho de arquivo (`max_file_size_mb`), para evitar
  que um PDF gigante travado num scanner mal configurado estoure a RAM do
  PC básico do usuário durante a extração de texto.
- Adicionado logging estruturado em vez de silenciar exceções por completo.
"""

import logging
import os
from pathlib import Path
from typing import List, TypedDict

from docx import Document
from pypdf import PdfReader

logger = logging.getLogger(__name__)


class ExtractedDocument(TypedDict):
    source: str
    content: str


class SafeDocumentReader:
    def __init__(
        self,
        allowed_paths: List[str],
        blocked_paths: List[str],
        max_file_size_mb: int = 25,
    ):
        self.allowed_paths = [Path(p).resolve() for p in allowed_paths]
        self.blocked_paths = [Path(p).resolve() for p in blocked_paths]
        self.max_file_size_bytes = max_file_size_mb * 1024 * 1024

    def _is_path_safe(self, file_path: Path) -> bool:
        """
        Um caminho só é considerado seguro se:
        1. Não estiver dentro de nenhuma pasta/arquivo bloqueado, E
        2. Estiver dentro de alguma pasta permitida.

        A checagem de bloqueio vem PRIMEIRO e tem prioridade: se um caminho
        está nas duas listas por engano, o bloqueio sempre vence
        (fail-safe: na dúvida, não expõe o dado).
        """
        resolved_file = file_path.resolve()

        for blocked in self.blocked_paths:
            if resolved_file == blocked or blocked in resolved_file.parents:
                return False

        for allowed in self.allowed_paths:
            if resolved_file == allowed or allowed in resolved_file.parents:
                return True

        return False

    def load_and_extract(self) -> List[ExtractedDocument]:
        extracted_documents: List[ExtractedDocument] = []

        for base_path in self.allowed_paths:
            if not base_path.exists():
                logger.warning("Caminho permitido não existe, ignorando: %s", base_path)
                continue

            for root, _, files in os.walk(base_path):
                for file in files:
                    file_path = Path(root) / file

                    if not self._is_path_safe(file_path):
                        continue

                    if not self._is_within_size_limit(file_path):
                        logger.warning(
                            "Arquivo ignorado por exceder %s MB: %s",
                            self.max_file_size_bytes // (1024 * 1024),
                            file_path,
                        )
                        continue

                    text = self._extract_text(file_path)

                    if text.strip():
                        extracted_documents.append({"source": str(file_path), "content": text})

        return extracted_documents

    def _is_within_size_limit(self, path: Path) -> bool:
        try:
            return path.stat().st_size <= self.max_file_size_bytes
        except OSError:
            return False

    def _extract_text(self, file_path: Path) -> str:
        suffix = file_path.suffix.lower()
        if suffix == ".pdf":
            return self._read_pdf(file_path)
        if suffix in (".docx", ".doc"):
            return self._read_docx(file_path)
        if suffix == ".txt":
            return self._read_txt(file_path)
        return ""

    def _read_pdf(self, path: Path) -> str:
        try:
            reader = PdfReader(str(path))
            return "\n".join(page.extract_text() or "" for page in reader.pages)
        except Exception:
            logger.exception("Falha ao ler PDF: %s", path)
            return ""

    def _read_docx(self, path: Path) -> str:
        try:
            doc = Document(str(path))
            return "\n".join(p.text for p in doc.paragraphs)
        except Exception:
            logger.exception("Falha ao ler DOCX: %s", path)
            return ""

    def _read_txt(self, path: Path) -> str:
        try:
            return path.read_text(encoding="utf-8", errors="ignore")
        except Exception:
            logger.exception("Falha ao ler TXT: %s", path)
            return ""
